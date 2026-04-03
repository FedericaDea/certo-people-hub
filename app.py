from core import *
from report import genera_report, esporta_kpi_excel
from openpyxl import load_workbook
import os

def mostra_richieste(nome_file):
    if not os.path.exists(nome_file):
        print("Nessuna richiesta trovata.")
        return

    wb = load_workbook(nome_file)
    ws = wb.active

    print("\n--- ELENCO RICHIESTE ---\n")

    for row in ws.iter_rows(min_row=2, values_only=True):
        print(f"ID: {row[0]} | Ruolo: {row[4]} | Stato: {row[9]} | Step: {row[10]}")

def aggiorna_stato(nome_file):
    if not os.path.exists(nome_file):
        print("File non trovato.")
        return

    id_da_modificare = input("Inserisci ID richiesta (es. RICH-001): ")

    wb = load_workbook(nome_file)
    ws = wb.active

    for row in ws.iter_rows(min_row=2):
        if row[0].value == id_da_modificare:

            print("1. Approvata")
            print("2. Rifiutata")
            print("3. Recruiting attivato")

            scelta = input("Seleziona nuovo stato: ")

            if scelta == "1":
                row[9].value = "Approvata"
                row[10].value = "Recruiting"
            elif scelta == "2":
                row[9].value = "Rifiutata"
                row[10].value = "-"
            elif scelta == "3":
                row[9].value = "Recruiting attivato"
                row[10].value = "HR"
            else:
                print("Scelta non valida")
                return

            wb.save(nome_file)
            print("Stato aggiornato!")
            return

    print("ID non trovato.")

def menu():
    nome_file_excel = "richieste_personale.xlsx"

    while True:
        print("\n--- CERT.O PEOPLE HUB ---")
        print("1. Crea richiesta")
        print("2. Visualizza richieste")
        print("3. Aggiorna stato richiesta")
        print("4. Genera report")
        print("5. Esporta KPI in Excel")
        print("6. Esci")

        scelta = input("Seleziona un'opzione: ")

        if scelta == "1":
            nome = input("Nome richiedente: ")
            reparto = input("Reparto: ")
            ruolo = input("Ruolo richiesto: ")
            motivazione = input("Motivazione: ")
            urgenza = input("Urgenza (Alta/Media/Bassa): ")
            tipo = input("Tipo richiesta: ")

            richiesta = crea_richiesta(
                nome, reparto, ruolo, motivazione, urgenza, tipo, nome_file_excel
            )

            salva_richiesta_excel(richiesta, nome_file_excel)

        elif scelta == "2":
            mostra_richieste(nome_file_excel)

        elif scelta == "3":
            aggiorna_stato(nome_file_excel)

        elif scelta == "4":
            genera_report(nome_file_excel)

        elif scelta == "5":
            esporta_kpi_excel(nome_file_excel)

        elif scelta == "6":
            print("Uscita dal sistema.")
            break

        else:
            print("Scelta non valida.")

menu()