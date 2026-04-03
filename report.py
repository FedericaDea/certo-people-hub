from openpyxl import load_workbook, Workbook
import os

def genera_report(nome_file):
    if not os.path.exists(nome_file):
        print("File richieste non trovato.")
        return

    wb = load_workbook(nome_file)
    ws = wb.active

    totale = 0
    approvate = 0
    rifiutate = 0
    in_validazione = 0
    recruiting_attivato = 0
    alta_priorita = 0

    reparti = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        totale += 1

        reparto = row[3]
        priorita = row[8]
        stato = row[9]

        if stato == "Approvata":
            approvate += 1
        elif stato == "Rifiutata":
            rifiutate += 1
        elif stato == "In validazione HR":
            in_validazione += 1
        elif stato == "Recruiting attivato":
            recruiting_attivato += 1

        if priorita == "PRIORITÀ ALTA":
            alta_priorita += 1

        if reparto in reparti:
            reparti[reparto] += 1
        else:
            reparti[reparto] = 1

    print("\n--- REPORT CERT.O PEOPLE HUB ---")
    print(f"Totale richieste: {totale}")
    print(f"Richieste approvate: {approvate}")
    print(f"Richieste rifiutate: {rifiutate}")
    print(f"Richieste in validazione HR: {in_validazione}")
    print(f"Recruiting attivato: {recruiting_attivato}")
    print(f"Richieste ad alta priorità: {alta_priorita}")

    print("\nRichieste per reparto:")
    for reparto, numero in reparti.items():
        print(f"- {reparto}: {numero}")


def esporta_kpi_excel(nome_file, nome_file_kpi="kpi_report.xlsx"):
    if not os.path.exists(nome_file):
        print("File richieste non trovato.")
        return

    wb = load_workbook(nome_file)
    ws = wb.active

    totale = 0
    approvate = 0
    rifiutate = 0
    in_validazione = 0
    recruiting_attivato = 0
    alta_priorita = 0

    reparti = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        totale += 1

        reparto = row[3]
        priorita = row[8]
        stato = row[9]

        if stato == "Approvata":
            approvate += 1
        elif stato == "Rifiutata":
            rifiutate += 1
        elif stato == "In validazione HR":
            in_validazione += 1
        elif stato == "Recruiting attivato":
            recruiting_attivato += 1

        if priorita == "PRIORITÀ ALTA":
            alta_priorita += 1

        if reparto in reparti:
            reparti[reparto] += 1
        else:
            reparti[reparto] = 1

    if totale > 0:
        perc_approvate = round((approvate / totale) * 100, 1)
        perc_rifiutate = round((rifiutate / totale) * 100, 1)
        perc_alta_priorita = round((alta_priorita / totale) * 100, 1)
    else:
        perc_approvate = 0
        perc_rifiutate = 0
        perc_alta_priorita = 0

    wb_kpi = Workbook()
    ws_kpi = wb_kpi.active
    ws_kpi.title = "KPI"

    ws_kpi.append(["Indicatore", "Valore"])
    ws_kpi.append(["Totale richieste", totale])
    ws_kpi.append(["Richieste approvate", approvate])
    ws_kpi.append(["Richieste rifiutate", rifiutate])
    ws_kpi.append(["Richieste in validazione HR", in_validazione])
    ws_kpi.append(["Recruiting attivato", recruiting_attivato])
    ws_kpi.append(["Richieste ad alta priorità", alta_priorita])
    ws_kpi.append(["% richieste approvate", f"{perc_approvate}%"])
    ws_kpi.append(["% richieste rifiutate", f"{perc_rifiutate}%"])
    ws_kpi.append(["% alta priorità", f"{perc_alta_priorita}%"])

    ws_reparti = wb_kpi.create_sheet("Reparti")
    ws_reparti.append(["Reparto", "Numero richieste"])

    for reparto, numero in reparti.items():
        ws_reparti.append([reparto, numero])

    ws_insight = wb_kpi.create_sheet("Insight")
    ws_insight.append(["Analisi automatica"])

    if perc_alta_priorita > 50:
        ws_insight.append(["Troppe richieste ad alta priorità"])
    else:
        ws_insight.append(["Priorità sotto controllo"])

    if approvate < rifiutate:
        ws_insight.append(["Più rifiuti che approvazioni"])
    else:
        ws_insight.append(["Processo equilibrato"])

    wb_kpi.save(nome_file_kpi)
    print(f"KPI esportati in {nome_file_kpi}")