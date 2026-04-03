from datetime import datetime
from openpyxl import Workbook, load_workbook
import os

def genera_id_richiesta(nome_file):
    if os.path.exists(nome_file):
        wb = load_workbook(nome_file)
        ws = wb.active
        ultimo_id = ws.max_row - 1
        nuovo_numero = ultimo_id + 1
    else:
        nuovo_numero = 1

    return f"RICH-{nuovo_numero:03d}"

def valida_richiesta_certo(ruolo_richiesto, motivazione, urgenza):
    note = []

    if not motivazione.strip():
        note.append("Manca la motivazione della richiesta")

    if urgenza.lower() not in ["alta", "media", "bassa"]:
        note.append("Urgenza non valida")

    if not ruolo_richiesto.strip():
        note.append("Manca il ruolo richiesto")

    if len(motivazione.strip()) < 10:
        note.append("Motivazione troppo breve")

    if len(note) == 0:
        esito = "Pronta per HR"
    else:
        esito = "Da rivedere"

    return esito, " | ".join(note)

def determina_stato_prossimo_step(esito_certo):
    if esito_certo == "Pronta per HR":
        return "In validazione HR", "HR"
    else:
        return "Da rivedere", "Manager"

def crea_richiesta(nome_richiedente, reparto, ruolo_richiesto, motivazione, urgenza, tipo_richiesta, nome_file):
    id_richiesta = genera_id_richiesta(nome_file)
    esito_certo, note_certo = valida_richiesta_certo(ruolo_richiesto, motivazione, urgenza)
    stato, prossimo_step = determina_stato_prossimo_step(esito_certo)

    richiesta = {
        "id_richiesta": id_richiesta,
        "data_richiesta": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "nome_richiedente": nome_richiedente,
        "reparto": reparto,
        "ruolo_richiesto": ruolo_richiesto,
        "motivazione": motivazione,
        "urgenza": urgenza,
        "tipo_richiesta": tipo_richiesta,
        "priorita": "PRIORITÀ ALTA" if urgenza.lower() == "alta" else "STANDARD",
        "stato": stato,
        "prossimo_step": prossimo_step,
        "esito_certo": esito_certo,
        "note_certo": note_certo
    }

    return richiesta

def salva_richiesta_excel(richiesta, nome_file="richieste_personale.xlsx"):
    if os.path.exists(nome_file):
        wb = load_workbook(nome_file)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Richieste"
        ws.append([
            "ID Richiesta",
            "Data Richiesta",
            "Nome Richiedente",
            "Reparto",
            "Ruolo Richiesto",
            "Motivazione",
            "Urgenza",
            "Tipo Richiesta",
            "Priorità",
            "Stato",
            "Prossimo Step",
            "Esito CERT.O",
            "Note CERT.O"
        ])

    ws.append([
        richiesta["id_richiesta"],
        richiesta["data_richiesta"],
        richiesta["nome_richiedente"],
        richiesta["reparto"],
        richiesta["ruolo_richiesto"],
        richiesta["motivazione"],
        richiesta["urgenza"],
        richiesta["tipo_richiesta"],
        richiesta["priorita"],
        richiesta["stato"],
        richiesta["prossimo_step"],
        richiesta["esito_certo"],
        richiesta["note_certo"]
    ])

    wb.save(nome_file)
    print(f"Richiesta salvata in {nome_file}")

nome_file_excel = "richieste_personale.xlsx"

richiesta1 = crea_richiesta(
    nome_richiedente="Mario Rossi",
    reparto="Sala",
    ruolo_richiesto="Cameriere",
    motivazione="Ampliamento staff per aumento coperti",
    urgenza="Alta",
    tipo_richiesta="Ampliamento staff",
    nome_file=nome_file_excel
)

salva_richiesta_excel(richiesta1, nome_file_excel)

print("ID richiesta:", richiesta1["id_richiesta"])
print("Stato richiesta:", richiesta1["stato"])
print("Prossimo step:", richiesta1["prossimo_step"])
print("Esito CERT.O:", richiesta1["esito_certo"])